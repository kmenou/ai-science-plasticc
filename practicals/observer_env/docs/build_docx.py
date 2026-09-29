"""Build the student guide as .docx from GUIDE.md plus the figures in docs/figures/.

    python docs/figures.py && python docs/build_docx.py      # run from practicals/observer_env/ (needs python-docx)
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SRC = Path("GUIDE.md")
OUT = Path("docs/Telescope_Eval_Student_Guide.docx")
FIG = Path("docs/figures")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x1F, 0x2A, 0x37)
CODE_INK = RGBColor(0x9A, 0x34, 0x12)
GREY = RGBColor(0x6B, 0x72, 0x80)
BODY_FONT, CODE_FONT = "Calibri", "Consolas"

# (anchor text that starts a paragraph, "before" | "after", figure file, width in inches, caption)
FIGURES = [
    ("The same environment can be driven by Claude", "before", "fig1_gym_loop.png", 6.2,
     "Figure 1. The Gym loop. Agent and environment exchange actions and observations until the agent submits; "
     "then the finished episode is graded."),
    ("We never write this loop ourselves.", "after", "fig2_architecture.png", 6.3,
     "Figure 2. What runs where. Our Python process owns the environment, the tools and the grader. The Claude "
     "Code subprocess runs the agent loop and talks to the model. The agent's only way into the data is the "
     "tool server."),
    ("Things to notice:", "before", "fig3_episode_dev006.png", 6.4,
     "Figure 3. Episode dev-006: telescope hours left after each tool call, for Claude Haiku (orange) and the "
     "heuristic script (blue). The red points mark the photo-z purchase, the decisive clue."),
    ("`env.finalize()` turns an episode", "before", "fig4_pipeline.png", 6.4,
     "Figure 4. From tasks to report: every episode leaves three records, and each is graded on output and "
     "process."),
    ("The cheater shows why output metrics", "before", "fig5_results.png", 5.2,
     "Figure 5. Pilot results on the dev set. The cheater is perfect on output alone; Claude Haiku matches the "
     "heuristic's accuracy but spends four times the telescope time."),
]


def shade(el, hex_fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def cell_borders(table, color="D1D5DB", size=4, left_accent=None):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        if left_accent and edge == "left":
            b.set(qn("w:sz"), "18")
            b.set(qn("w:color"), left_accent)
        borders.append(b)
    tbl.tblPr.append(borders)


def add_inline(par, text, size=None, color=None):
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            add_inline_bold(par, tok[2:-2], size, color)
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = CODE_INK
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1])
            r.italic = True
            _style(r, size, color)
        else:
            _style(par.add_run(tok), size, color)


def add_inline_bold(par, text, size, color):
    for tok in re.split(r"(`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = CODE_INK
        else:
            r = par.add_run(tok)
            _style(r, size, color)
        r.bold = True


def _style(run, size, color):
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def code_block(doc, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_borders(t, color="E5E7EB", left_accent="2A78D6")
    c = t.cell(0, 0)
    shade(c._tc, "F4F6F8")
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    longest = max(len(l) for l in lines)
    size = min(8.5, 455 / (0.6 * longest))          # monospace ~0.6 em per char; ~455 pt usable width
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = CODE_FONT
        r.font.size = Pt(round(size, 1))
        r.font.color.rgb = INK
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def md_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_borders(t)
    for j, h in enumerate(header):
        c = t.cell(0, j)
        shade(c._tc, "DCE7F5")
        p = c.paragraphs[0]
        add_inline_bold(p, h, 9.5, NAVY)
    for i, row in enumerate(body, 1):
        for j, val in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            add_inline(p, val, 9.5)
            if i % 2 == 0:
                shade(t.cell(i, j)._tc, "F9FAFB")
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, fname, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(FIG / fname), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(10)
    r = cp.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)
    r.font.size = Pt(9)
    r.font.color.rgb = GREY


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.9))
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    page_number_footer(sec)
    st = doc.styles
    n = st["Normal"]
    n.font.name = BODY_FONT
    n.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    n.font.size = Pt(10.5)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.12
    for name, size, before in (("Title", 24, 0), ("Heading 1", 15, 16)):
        s = st[name]
        s.font.name = BODY_FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = NAVY
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2"):
        st[name].font.name = BODY_FONT
        st[name].font.size = Pt(10.5)
        st[name].paragraph_format.space_after = Pt(3)


def build():
    md = SRC.read_text()
    doc = Document()
    setup(doc)
    lines = md.splitlines()
    i = 0

    def flush_para(text):
        text = text.strip()
        if not text:
            return
        for anchor, where, f, w, cap in FIGURES:
            if text.startswith(anchor) and where == "before":
                figure(doc, f, w, cap)
        p = doc.add_paragraph()
        add_inline(p, text)
        for anchor, where, f, w, cap in FIGURES:
            if text.startswith(anchor) and where == "after":
                figure(doc, f, w, cap)

    para: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush_para(" ".join(para)); para = []
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            code_block(doc, lines[i + 1:j])
            i = j + 1
            continue
        if line.startswith("|"):
            flush_para(" ".join(para)); para = []
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            md_table(doc, lines[i:j])
            i = j
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title")
            sub = doc.add_paragraph()
            r = sub.add_run("PLAsTiCC course  ·  agentic evaluation practical  ·  student guide")
            r.font.color.rgb = GREY
            r.font.size = Pt(11)
            i += 1
            continue
        if line.startswith("## "):
            flush_para(" ".join(para)); para = []
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
            i += 1
            continue
        if line.strip() == "---":
            flush_para(" ".join(para)); para = []
            i += 1
            continue
        m_b = re.match(r"^(\s*)- (.*)", line)
        m_n = re.match(r"^(\d+)\. (.*)", line)
        if m_b or m_n:
            flush_para(" ".join(para)); para = []
            text = (m_b.group(2) if m_b else m_n.group(2))
            j = i + 1
            while j < len(lines) and lines[j].startswith("  ") and not re.match(r"^\s*- |^\d+\. ", lines[j]) \
                    and lines[j].strip():
                text += " " + lines[j].strip()
                j += 1
            if m_b:
                style = "List Bullet 2" if len(m_b.group(1)) >= 2 else "List Bullet"
                add_inline(doc.add_paragraph(style=style), text)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.22)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(f"{m_n.group(1)}.  ")
                r.bold = True
                r.font.color.rgb = NAVY
                add_inline(p, text)
            i = j
            continue
        if not line.strip():
            flush_para(" ".join(para)); para = []
            i += 1
            continue
        para.append(line.strip())
        i += 1
    flush_para(" ".join(para))
    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    print("->", OUT)


if __name__ == "__main__":
    build()
